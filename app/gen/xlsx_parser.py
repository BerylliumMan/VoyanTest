"""XLSX 解析：把工作表内容转为纯文本（供生成管线 / 对话附件使用）。

- 只读模式 + data_only（读缓存值而非公式）；
- 每表行数上限与总字符上限，超限显式标注（不静默截断）。
"""
from __future__ import annotations

import io
import logging

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

MAX_ROWS_PER_SHEET = 500
MAX_CHARS = 200_000


def extract_text_from_xlsx(file, *, max_rows: int = MAX_ROWS_PER_SHEET) -> str:
    """把 xlsx 文件（file-like 或 bytes）转为文本。

    输出形如：
        ===== 工作表：Sheet1 =====
        列1\t列2\t列3
        ...
    空行跳过；超出 max_rows 的行截断并标注。
    """
    if hasattr(file, "read"):
        file.seek(0)
        data = file.read()
    else:
        data = file
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    total_chars = 0
    try:
        for ws in wb.worksheets:
            lines = [f"===== 工作表：{ws.title} ====="]
            rows = 0
            for row in ws.iter_rows(values_only=True):
                if row is None:
                    continue
                cells = ["" if v is None else str(v) for v in row]
                if not any(c.strip() for c in cells):
                    continue
                lines.append("\t".join(cells).rstrip())
                rows += 1
                if rows >= max_rows:
                    lines.append(f"…（本工作表超过 {max_rows} 行，已截断）")
                    break
            part = "\n".join(lines)
            parts.append(part)
            total_chars += len(part)
            if total_chars > MAX_CHARS:
                parts.append("…（内容过长，已截断）")
                break
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            logger.debug("xlsx close failed", exc_info=True)
    return "\n\n".join(parts)


__all__ = ["extract_text_from_xlsx"]
