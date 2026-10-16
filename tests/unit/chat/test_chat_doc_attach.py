"""035+：文档附件能力单测 —— xlsx 解析 / 附件内容提取 / 文档驱动探索接线。

覆盖：
- chat_files.resolve_chat_file 路径守卫（穿越/绝对路径/缺失）
- extract_attachment_text：xlsx / docx / md / 截断 / 图片说明 / 不支持类型
- contracts：browse_generate.document_ref 透传
- context：file ref 的文本注入渲染（text / text_error）
- 工具面：generate_cases.file_ref、browse_site.document_ref 参数存在
- orchestrator：发送消息时附件内容提取并写入 refs
"""
import io
import os

import pytest

from app.services import chat_files
from core.chat import context as chat_context
from core.chat import contracts as chat_contracts


def _make_xlsx(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    ws.append(["编号", "步骤", "预期"])
    ws.append(["TC001", "打开登录页", "显示登录表单"])
    ws.append(["TC002", "输入错误密码", "提示用户名或密码错误"])
    wb.save(path)


def _make_docx(path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("测试说明：验证登录流程")
    doc.add_paragraph("步骤1：打开登录页并确认表单显示")
    doc.save(path)


# ----------------------------
# resolve 守卫
# ----------------------------

def test_resolve_chat_file_guard(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    os.makedirs("data/chat_uploads/x", exist_ok=True)
    with open("data/chat_uploads/x/a.md", "w") as fh:
        fh.write("hi")
    assert chat_files.resolve_chat_file("data/chat_uploads/x/a.md") is not None
    assert chat_files.resolve_chat_file("data/chat_uploads/../secret.md") is None  # 路径穿越
    assert chat_files.resolve_chat_file("/etc/passwd") is None
    assert chat_files.resolve_chat_file("data/chat_uploads/x/missing.md") is None
    assert chat_files.resolve_chat_file("") is None


# ----------------------------
# 文本提取
# ----------------------------

def test_extract_xlsx_roundtrip(tmp_path):
    p = tmp_path / "cases.xlsx"
    _make_xlsx(str(p))
    text, err = chat_files.extract_attachment_text(str(p))
    assert err == ""
    assert "工作表：测试用例" in text
    assert "TC001" in text and "打开登录页" in text and "显示登录表单" in text
    assert "TC002" in text


def test_extract_docx_roundtrip(tmp_path):
    p = tmp_path / "spec.docx"
    _make_docx(str(p))
    text, err = chat_files.extract_attachment_text(str(p))
    assert err == ""
    assert "测试说明" in text and "步骤1" in text


def test_extract_md_and_truncation(tmp_path):
    p = tmp_path / "notes.md"
    p.write_text("A" * 10000, encoding="utf-8")
    text, err = chat_files.extract_attachment_text(str(p), max_chars=500)
    assert err == ""
    assert len(text) <= 500 + 30  # 截断标注少量余量
    assert "已截断" in text


def test_extract_image_returns_note(tmp_path):
    p = tmp_path / "shot.png"
    p.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
    text, err = chat_files.extract_attachment_text(str(p))
    assert text == ""
    assert "图片附件" in err


def test_extract_unsupported_ext(tmp_path):
    p = tmp_path / "x.bin"
    p.write_bytes(b"\x00\x01")
    text, err = chat_files.extract_attachment_text(str(p))
    assert text == "" and "不支持" in err


# ----------------------------
# 契约：document_ref 透传
# ----------------------------

def test_contract_browse_document_ref_passthrough():
    norm, err = chat_contracts.validate_action({
        "type": "browse_generate",
        "url": "http://localhost:8002",
        "goal": "按文档执行测试",
        "document_ref": "data/chat_uploads/abc/spec.docx",
    })
    assert err == ""
    assert norm["document_ref"] == "data/chat_uploads/abc/spec.docx"


def test_contract_generate_file_source_passthrough():
    norm, err = chat_contracts.validate_action({
        "type": "generate", "kind": "ui", "source": "file",
        "payload": {"file_ref": "data/chat_uploads/abc/spec.xlsx"},
    })
    assert err == ""
    assert norm["source"] == "file"
    assert norm["payload"]["file_ref"] == "data/chat_uploads/abc/spec.xlsx"


# ----------------------------
# 上下文渲染
# ----------------------------

def test_context_renders_file_text():
    msg = {
        "role": "user",
        "content": "按附件生成用例",
        "refs": [{
            "type": "file",
            "file_ref": "data/chat_uploads/x/a.xlsx",
            "label": "a.xlsx",
            "text": "工作表：测试用例\nTC001\t打开登录页",
        }],
    }
    out = chat_context.message_text(msg)
    assert "【文件内容（节选）】" in out
    assert "TC001" in out


def test_context_renders_file_text_error():
    msg = {
        "role": "user",
        "content": "看附件",
        "refs": [{
            "type": "file",
            "file_ref": "data/chat_uploads/x/a.png",
            "label": "a.png",
            "text_error": "图片附件（未提取文本）",
        }],
    }
    out = chat_context.message_text(msg)
    assert "文件未能提取内容" in out


# ----------------------------
# 工具面
# ----------------------------

def test_tool_specs_have_document_params():
    import core.tools.builtin.actions  # noqa: F401  触发注册
    from core.tools.registry import REGISTRY

    gen = REGISTRY.get("generate_cases")
    assert gen is not None
    assert "file_ref" in (gen.parameters.get("properties") or {})

    browse = REGISTRY.get("browse_site")
    assert browse is not None
    assert "document_ref" in (browse.parameters.get("properties") or {})


# ----------------------------
# orchestrator：发送时提取附件内容进 refs
# ----------------------------

@pytest.mark.asyncio
async def test_handle_user_message_extracts_attachment(tmp_path, monkeypatch, chat_store, fake_llm_factory):
    from core.chat import orchestrator as orch
    from tests.unit.chat.conftest import FakeConversation

    p = tmp_path / "cases.xlsx"
    _make_xlsx(str(p))
    monkeypatch.setattr(chat_files, "resolve_chat_file", lambda ref: str(p))

    conv = FakeConversation()
    llm = fake_llm_factory({"reply": "收到文档", "actions": []})
    user_msg, _ = await orch.handle_user_message(
        None, conversation=conv, user_id=7, content="根据附件生成用例",
        attachments=[{"file_ref": "data/chat_uploads/x/cases.xlsx", "filename": "cases.xlsx"}],
        llm_call=llm,
    )
    refs = user_msg.refs or []
    file_ref = next(r for r in refs if r.get("type") == "file")
    assert "TC001" in (file_ref.get("text") or "")
    # 模型上下文包含文件内容
    sent = llm.calls[0]
    assert any("TC001" in m["content"] for m in sent)
